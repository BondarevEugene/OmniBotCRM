"""Spreadsheet exports with safe text cells and a consistent readable layout."""
from io import BytesIO
from datetime import datetime
from decimal import Decimal

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
        return "'" + value
    return value


def workbook_response(filename: str, sheets):
    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for row in rows:
            ws.append([_cell(v) for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='303F73')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 24
        for column in ws.columns:
            col = list(column)
            width = min(52, max(12, max(len(str(c.value or '')) for c in col[:500]) + 2))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    out = BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename="{filename}"',
                             'Cache-Control': 'no-store', 'X-Content-Type-Options': 'nosniff'})
