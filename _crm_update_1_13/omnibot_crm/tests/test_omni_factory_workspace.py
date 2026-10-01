"""Cross-module checks for no-code bot workspace, ads, payments and Excel."""
import base64
import hashlib
import hmac
import json
import os
import re
import time
from datetime import date
from decimal import Decimal
from io import BytesIO

os.environ.setdefault('SESSION_SECRET', 'test-secret-with-at-least-32-characters')
os.environ.setdefault('COOKIE_SECURE', 'false')

import httpx
import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from crm.acquirers import save_acquirer
from crm.ad_reporting import sync_account
from crm.app import create_app
from crm.config import Settings
from crm.db import Base
from crm.models import (Admin, Bot, BotPaymentProvider, BotScenario, Customer,
                        MarketingAdAccount, MarketingCampaign, MarketingChannel,
                        MarketingMetric, Order, Payment)
from crm.payment_providers import VerificationError, verify_payment
from crm.security import hash_password


def login(client):
    page = client.get('/login')
    token = re.search(r'name="token" value="([^"]+)', page.text).group(1)
    assert client.post('/login', data={'token': token, 'email': 'owner@example.com',
        'password': 'strong-password-123'}, follow_redirects=False).status_code == 303
    return token


def test_bot_creation_scenario_gateway_and_excel(tmp_path):
    secret = 'a-long-unique-event-secret-for-my-bot'
    config = Settings(database_url=f'sqlite:///{tmp_path / "crm.db"}',
        session_secret='s' * 40, cookie_secure=False, encryption_key=Fernet.generate_key().decode(),
        ingest_secrets={'new_bot': secret})
    app = create_app(config)
    Base.metadata.create_all(app.state.engine)
    with app.state.session_factory() as db:
        db.add(Admin(email='owner@example.com', password_hash=hash_password('strong-password-123')))
        db.commit()
    with TestClient(app) as client:
        token = login(client)
        assert client.get('/marketing').status_code == 200
        bots_page = client.get('/bots').text
        assert 'Добавить бота' in bots_page
        token = re.search(r'name="token" value="([^"]+)', bots_page).group(1)
        created = client.post('/bots', data={'token': token, 'key': 'new_bot', 'name': 'Новый бот'},
                              follow_redirects=False)
        assert created.status_code == 303 and created.headers['location'].endswith('/configure#modules')
        assert client.post('/bots/1/scenarios', data={'token': token, 'key': 'greeting',
            'name': 'Приветствие', 'trigger_kind': 'command', 'trigger_value': '/start'},
            follow_redirects=False).status_code == 303
        assert client.post('/bots/1/scenarios/1/steps', data={'token': token, 'kind': 'question',
            'text': 'Какая задача?', 'field_key': 'project_need', 'choices': 'Сайт, Бот'},
            follow_redirects=False).status_code == 303
        stamp = str(int(time.time()))
        signature = hmac.new(secret.encode(), (stamp + '.GET./api/v1/bot-config').encode(),
                             hashlib.sha256).hexdigest()
        headers = {'X-Bot-Key': 'new_bot', 'X-Timestamp': stamp, 'X-Signature': signature}
        draft = client.get('/api/v1/bot-config', headers=headers)
        assert draft.status_code == 200 and draft.json()['scenarios'] == []
        assert client.post('/bots/1/scenarios/1/publish', data={'token': token, 'status': 'published'},
            follow_redirects=False).status_code == 303
        published = client.get('/api/v1/bot-config', headers=headers)
        assert published.json()['scenarios'][0]['steps'][0]['options']['field'] == 'project_need'
        raw = json.dumps({'event_id': 'flow-1', 'type': 'scenario.enter', 'data': {
            'external_user_id': 'person-1', 'scenario_key': 'greeting'}}).encode()
        event_sig = hmac.new(secret.encode(), stamp.encode() + b'.' + raw,
                             hashlib.sha256).hexdigest()
        assert client.post('/api/v1/events', content=raw, headers={**headers,
            'X-Signature': event_sig}).status_code == 200
        detail = client.get('/bots/1').text
        assert '1 входов' in detail
        token = re.search(r'name="token" value="([^"]+)', detail).group(1)
        assert client.get('/api/v1/bot-config', headers={**headers,
            'If-None-Match': published.headers['ETag']}).status_code == 304
        assert client.post('/bots/1/scenarios/1/steps/1/order', data={'token': token, 'action': 'remove'},
            follow_redirects=False).status_code == 303
        assert client.get('/api/v1/bot-config', headers=headers).json()['scenarios'] == []
        # A regular valid CSV import with duplicate update and explicit no-consent behavior.
        payload = 'name,email,phone\nAlice,alice@example.com,+380000000001\n'
        assert client.post('/customers/import', data={'token': token},
            files={'upload': ('contacts.csv', payload, 'text/csv')},
            follow_redirects=False).status_code == 303
        assert client.post('/customers/import', data={'token': token},
            files={'upload': ('contacts.csv', payload, 'text/csv')},
            follow_redirects=False).status_code == 303
        sheet = client.get('/customers/export.xlsx')
        assert sheet.status_code == 200
        wb = load_workbook(BytesIO(sheet.content), read_only=True)
        assert wb.sheetnames == ['Клиенты', 'Боты']
        assert 'Alice' in [row[1] for row in list(wb['Клиенты'].values)[1:]]
        assert client.get('/m/export/bot_settings.xlsx').status_code == 200
        assert client.get('/bots/1/configure').status_code == 200
        with app.state.session_factory() as db:
            assert db.scalar(select(BotScenario)).status == 'draft'
            assert len(db.scalars(select(Customer)).all()) == 2
            assert not any(c.marketing_consent for c in db.scalars(select(Customer)))
    app.state.engine.dispose()


def test_ads_meta_google_daily_import_is_isolated_and_idempotent(tmp_path):
    config = Settings(database_url=f'sqlite:///{tmp_path / "crm.db"}', session_secret='s'*40,
                      cookie_secure=False, encryption_key=Fernet.generate_key().decode())
    app = create_app(config)
    Base.metadata.create_all(app.state.engine)
    with app.state.session_factory() as db:
        meta = MarketingChannel(platform='facebook', account='Meta')
        google = MarketingChannel(platform='google_ads', account='Google')
        db.add_all([meta, google, Admin(email='owner@example.com',
            password_hash=hash_password('strong-password-123'))]); db.flush()
        from crm.bot_settings import encrypt
        db.add_all([
            MarketingAdAccount(channel_id=meta.id, external_id='12345', currency='USD',
                api_version='v25.0', enabled=True,
                encrypted_credentials=encrypt(json.dumps({'access_token': 'test'}), config.encryption_key)),
            MarketingAdAccount(channel_id=google.id, external_id='1234567890', currency='EUR',
                api_version='v25', enabled=True,
                encrypted_credentials=encrypt(json.dumps({'client_id': 'a', 'client_secret': 'b',
                    'refresh_token': 'c'}), config.encryption_key))])
        db.commit()
        from datetime import datetime, timezone, timedelta
        day = (datetime.now(timezone.utc).date() - timedelta(days=1)).isoformat()
        def handler(request):
            if request.url.host == 'graph.facebook.com':
                assert request.headers['authorization'] == 'Bearer test'
                return httpx.Response(200, json={'data': [{'campaign_id': '8', 'campaign_name': 'Campaign',
                    'date_start': day, 'impressions': '100', 'clicks': '7', 'spend': '12.50',
                    'actions': [{'action_type': 'lead', 'value': '2'}]}]})
            if request.url.host == 'oauth2.googleapis.com':
                return httpx.Response(200, json={'access_token': 'short-lived'})
            assert request.url.host == 'googleads.googleapis.com'
            return httpx.Response(200, json={'results': [{'campaign': {'id': '8', 'name': 'Campaign'},
                'segments': {'date': day}, 'metrics': {'impressions': '200', 'clicks': '8',
                    'costMicros': '5000000', 'conversions': '1.5'}}]})
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            for _ in range(2):
                with app.state.session_factory() as db:
                    for account in db.scalars(select(MarketingAdAccount).order_by(MarketingAdAccount.id)):
                        assert sync_account(db, account, config.encryption_key, client) == 1
                    db.commit()
        with app.state.session_factory() as db:
            campaigns = db.scalars(select(MarketingCampaign)).all()
            metrics = db.scalars(select(MarketingMetric)).all()
            assert len(campaigns) == len(metrics) == 2
            assert {m.campaign.currency for m in metrics} == {'USD', 'EUR'}
            assert sum(m.conversions for m in metrics) == Decimal('3.50')
            assert sum(m.leads for m in metrics) == 2
    with TestClient(app) as browser:
        login(browser)
        page = browser.get('/marketing')
        assert page.status_code == 200
        assert '1234567890' in page.text and 'Синхронизировано' in page.text
        assert browser.get('/marketing/export.xlsx').status_code == 200
    app.state.engine.dispose()


def test_liqpay_and_paypal_require_real_provider_proof(tmp_path):
    config = Settings(database_url=f'sqlite:///{tmp_path / "crm.db"}', session_secret='s'*40,
                      cookie_secure=False, encryption_key=Fernet.generate_key().decode())
    app = create_app(config)
    Base.metadata.create_all(app.state.engine)
    with app.state.session_factory() as db:
        bot = Bot(key='shop', name='Shop')
        customer = Customer(name='Customer')
        db.add_all([bot, customer]); db.flush()
        order_a = Order(bot_id=bot.id, customer_id=customer.id, external_id='order-liq',
                        title='A', amount=Decimal('7.00'), currency='UAH')
        order_b = Order(bot_id=bot.id, customer_id=customer.id, external_id='5O190127TN364715T',
                        title='B', amount=Decimal('9.00'), currency='USD')
        db.add_all([order_a, order_b]); db.flush()
        db.add_all([Payment(bot_id=bot.id, customer_id=customer.id, order_id=order_a.id,
            external_id='p-a', provider='liqpay', amount=Decimal('7.00'), currency='UAH'),
            Payment(bot_id=bot.id, customer_id=customer.id, order_id=order_b.id,
            external_id='p-b', provider='paypal', amount=Decimal('9.00'), currency='USD')])
        save_acquirer(db, bot.id, 'liqpay', '', {'public_key': 'public', 'private_key': 'private'},
                      '', True, config.encryption_key, None)
        save_acquirer(db, bot.id, 'paypal', '', {'client_id': 'client', 'client_secret': 'secret',
            'merchant_id': 'merchant', 'environment': 'live'}, '', True, config.encryption_key, None)
        db.commit()
        def handler(request):
            if request.url.host == 'www.liqpay.ua':
                return httpx.Response(200, json={'public_key': 'public', 'order_id': 'order-liq',
                    'result': 'ok', 'status': 'success', 'action': 'pay', 'amount': '7.00',
                    'currency': 'UAH', 'payment_id': 'liq-123'})
            if request.url.path == '/v1/oauth2/token':
                return httpx.Response(200, json={'access_token': 'access'})
            return httpx.Response(200, json={'id': '5O190127TN364715T', 'status': 'COMPLETED',
                'purchase_units': [{'payee': {'merchant_id': 'merchant'}, 'payments': {'captures': [
                    {'id': 'capture-123', 'status': 'COMPLETED',
                     'amount': {'value': '9.00', 'currency_code': 'USD'}}]}}]})
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            payments = db.scalars(select(Payment).order_by(Payment.id)).all()
            assert verify_payment(db, payments[0], config, client) == 'liq-123'
            assert verify_payment(db, payments[1], config, client) == 'capture-123'
            db.commit()
            assert [p.status for p in payments] == ['paid', 'paid']
        def wrong_merchant(request):
            if request.url.path == '/v1/oauth2/token':
                return httpx.Response(200, json={'access_token': 'access'})
            return httpx.Response(200, json={'id': '5O190127TN364715T', 'status': 'COMPLETED',
                'purchase_units': [{'payee': {'merchant_id': 'someone-else'}, 'payments': {
                    'captures': [{'id': 'capture-x', 'status': 'COMPLETED',
                    'amount': {'value': '9.00', 'currency_code': 'USD'}}]}}]})
        with httpx.Client(transport=httpx.MockTransport(wrong_merchant)) as client:
            with pytest.raises(VerificationError, match='merchant'):
                verify_payment(db, payments[1], config, client)
        save_acquirer(db, bot.id, 'paypal', '', {'environment': 'sandbox'},
                      '', True, config.encryption_key, None)
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            with pytest.raises(VerificationError, match='sandbox'):
                verify_payment(db, payments[1], config, client)
    app.state.engine.dispose()
